import csv
import re
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_EXTRACTED_CHARS = 250_000

CATEGORY_RULES = {
    "Personal Data": ["name", "email", "phone", "address", "customer", "account", "profile"],
    "Sensitive Data": ["national id", "passport", "health", "medical", "biometric", "credit card", "card number", "iban", "bank account"],
    "Usage Data": ["ip address", "user agent", "cookie", "browser", "device id", "session"],
    "Support Data": ["support ticket", "ticket", "case number", "attachment", "support request"],
    "Financial Data": ["invoice", "payment", "transaction", "credit card", "iban", "bank account"],
    "Employee Data": ["employee", "payroll", "salary", "human resources", "hr record"],
    "Business Confidential Data": ["confidential", "contract", "revenue", "budget", "internal project"],
    "Technical Security Data": ["api key", "secret", "token", "password", "firewall", "log", "security event"],
}

def _cap(value: str) -> str:
    return value[:MAX_EXTRACTED_CHARS]

def _documents_control(text: str, terms: list[str]) -> bool:
    """Return True only when a control is positively documented, not negated."""
    negative_context = re.compile(
        r"\b(?:no|not|without|lack|lacks|missing|absent|does not|do not|is not|isn't|under review)\b"
    )
    for term in terms:
        for match in re.finditer(re.escape(term), text):
            window = text[max(0, match.start() - 70):match.end() + 70]
            if not negative_context.search(window):
                return True
    return False

def extract_text(file_path: Path, suffix: str) -> str:
    suffix = suffix.lower()

    if suffix == ".txt":
        return _cap(file_path.read_text(encoding="utf-8", errors="ignore"))

    if suffix == ".csv":
        rows = []
        with file_path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
            for row in csv.reader(handle):
                rows.append(" | ".join(str(cell) for cell in row))
                if len("\n".join(rows)) >= MAX_EXTRACTED_CHARS:
                    break
        return _cap("\n".join(rows))

    if suffix == ".docx":
        document = Document(file_path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return _cap("\n".join(parts))

    if suffix == ".xlsx":
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    parts.append(" | ".join(values))
                if len("\n".join(parts)) >= MAX_EXTRACTED_CHARS:
                    return _cap("\n".join(parts))
        return _cap("\n".join(parts))

    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
            if len("\n".join(parts)) >= MAX_EXTRACTED_CHARS:
                break
        return _cap("\n".join(parts))

    raise ValueError(f"Unsupported file type: {suffix}")

def analyze_text(text: str, processors: list[str] | None = None) -> dict:
    normalized = text.lower()
    categories = []
    for category, keywords in CATEGORY_RULES.items():
        if any(keyword in normalized for keyword in keywords):
            categories.append(category)

    if not categories:
        categories = ["Technical Security Data"] if normalized.strip() else []

    sensitive = "Sensitive Data" in categories
    financial = "Financial Data" in categories

    # Generic terms such as "account" appear in firewall logs and telemetry.
    # Treat them as PII only when the content also contains stronger identity context.
    strong_personal_markers = (
        "email address", "phone number", "full name", "first name", "last name",
        "customer name", "home address", "mailing address", "date of birth",
        "account holder", "employee record", "hr record", "payroll", "salary",
    )
    has_strong_personal_context = any(marker in normalized for marker in strong_personal_markers)
    technical_or_usage_only = (
        ("Technical Security Data" in categories or "Usage Data" in categories)
        and not has_strong_personal_context
        and "Employee Data" not in categories
        and "Sensitive Data" not in categories
        and "Financial Data" not in categories
        and "Business Confidential Data" not in categories
    )

    personal = (
        ("Employee Data" in categories)
        or ("Personal Data" in categories and has_strong_personal_context)
    ) and not technical_or_usage_only

    if sensitive:
        classification = "Sensitive PII"
    elif personal and ("account" in normalized or "phone" in normalized or "address" in normalized or financial):
        classification = "Restricted PII"
    elif personal or financial or "Business Confidential Data" in categories or "Employee Data" in categories:
        classification = "Confidential"
    elif categories:
        classification = "Internal"
    else:
        classification = "Unknown"

    controls = set()
    if personal or sensitive or financial:
        controls.update(["gdpr-article-5", "gdpr-article-25", "gdpr-article-32", "iso-a-5-34"])
    if sensitive or financial:
        controls.update(["iso-a-8-12", "iso-a-8-24"])
    if "Technical Security Data" in categories:
        controls.add("iso-a-8-15")
    if processors:
        controls.add("gdpr-article-28")

    findings = []
    has_personal_context = personal or sensitive or financial
    mentions_mfa = _documents_control(normalized, ["mfa", "multi-factor", "multifactor"])
    mentions_retention = _documents_control(normalized, ["retention", "deletion schedule", "delete after"])
    mentions_encryption = _documents_control(normalized, ["encrypt", "tls"])

    if has_personal_context and not mentions_mfa:
        findings.append({
            "title": "Privileged access lacks MFA",
            "rationale": "The assessed content indicates personal or sensitive data processing but does not document MFA for privileged access.",
            "likelihood": "High", "impact": "High", "score": 20,
            "controls": ["gdpr-article-32", "iso-a-5-15"],
        })
    if has_personal_context and not mentions_retention:
        findings.append({
            "title": "Retention period not defined",
            "rationale": "The assessment does not identify a retention or deletion schedule for processed data.",
            "likelihood": "Medium", "impact": "High", "score": 16,
            "controls": ["gdpr-article-5", "iso-a-8-10"],
        })
    if processors:
        findings.append({
            "title": "Third-party processor review required",
            "rationale": "A third-party processor is linked to this asset and requires appropriate processing governance.",
            "likelihood": "Medium", "impact": "Medium", "score": 12,
            "controls": ["gdpr-article-28", "iso-a-5-34"],
        })
    if has_personal_context and not mentions_encryption:
        findings.append({
            "title": "Encryption control requires validation",
            "rationale": "The assessment does not document encryption for personal or sensitive data in transit and at rest.",
            "likelihood": "Medium", "impact": "High", "score": 16,
            "controls": ["gdpr-article-32", "iso-a-8-24"],
        })

    risk_score = max((item["score"] for item in findings), default=4)
    risk_level = "High" if risk_score >= 16 else "Medium" if risk_score >= 8 else "Low"

    for finding in findings:
        controls.update(finding["controls"])

    return {
        "categories": categories,
        "classification": classification,
        "risk_score": risk_score,
        "risk_level": risk_level,
        "control_ids": sorted(controls),
        "findings": findings,
    }
